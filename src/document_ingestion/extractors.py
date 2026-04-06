from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List

from .models import ExtractedDocument, ExtractedPage


class UnsupportedDocumentError(ValueError):
    """Raised when the uploaded document type is not supported."""


def normalize_extracted_text(text: str) -> str:
    raw = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    paragraphs: List[str] = []
    current: List[str] = []
    for line in raw.split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if not cleaned:
            if current:
                paragraphs.append(" ".join(current).strip())
                current = []
            continue
        current.append(cleaned)
    if current:
        paragraphs.append(" ".join(current).strip())
    return "\n\n".join(paragraph for paragraph in paragraphs if paragraph)


class PdfTextExtractor:
    def __init__(self, min_text_chars: int = 80):
        self.min_text_chars = int(min_text_chars)

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        import fitz

        path = Path(file_path)
        document = fitz.open(path)
        pages: List[ExtractedPage] = []
        all_pages: List[str] = []
        cursor = 0

        for index in range(len(document)):
            page = document.load_page(index)
            page_text = page.get_text("text") or ""
            block_entries = page.get_text("blocks") or []
            blocks = [
                normalize_extracted_text(block[4])
                for block in block_entries
                if len(block) >= 5 and normalize_extracted_text(block[4])
            ]
            normalized_page = normalize_extracted_text(page_text)
            start_char = cursor
            end_char = start_char + len(normalized_page)
            pages.append(
                ExtractedPage(
                    page_number=index + 1,
                    text=normalized_page,
                    blocks=blocks,
                    start_char=start_char,
                    end_char=end_char,
                )
            )
            if normalized_page:
                all_pages.append(normalized_page)
                cursor = end_char + 2

        document.close()
        normalized_text = "\n\n".join(part for part in all_pages if part)
        visible_chars = len(re.sub(r"\W+", "", normalized_text, flags=re.UNICODE))
        needs_ocr = visible_chars < self.min_text_chars
        return ExtractedDocument(
            file_path=path,
            file_type="pdf",
            parser_used="pymupdf",
            text=normalized_text,
            normalized_text=normalized_text,
            pages=pages,
            needs_ocr=needs_ocr,
        )


class DocxTextExtractor:
    def extract(self, file_path: str | Path) -> ExtractedDocument:
        from docx import Document

        path = Path(file_path)
        document = Document(path)
        paragraphs = [normalize_extracted_text(paragraph.text) for paragraph in document.paragraphs]
        paragraphs = [paragraph for paragraph in paragraphs if paragraph]
        normalized_text = "\n\n".join(paragraphs)
        page = ExtractedPage(
            page_number=1,
            text=normalized_text,
            blocks=paragraphs,
            start_char=0,
            end_char=len(normalized_text),
        )
        return ExtractedDocument(
            file_path=path,
            file_type="docx",
            parser_used="python-docx",
            text=normalized_text,
            normalized_text=normalized_text,
            pages=[page] if normalized_text else [],
        )


class TxtTextExtractor:
    ENCODINGS = ("utf-8", "utf-8-sig", "cp1258", "latin-1")

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        path = Path(file_path)
        raw_bytes = path.read_bytes()
        decoded = None
        for encoding in self.ENCODINGS:
            try:
                decoded = raw_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            decoded = raw_bytes.decode("utf-8", errors="ignore")

        normalized_text = normalize_extracted_text(decoded)
        page = ExtractedPage(
            page_number=1,
            text=normalized_text,
            blocks=[paragraph for paragraph in normalized_text.split("\n\n") if paragraph],
            start_char=0,
            end_char=len(normalized_text),
        )
        return ExtractedDocument(
            file_path=path,
            file_type="txt",
            parser_used="plain-text",
            text=decoded,
            normalized_text=normalized_text,
            pages=[page] if normalized_text else [],
        )


class ExtractorRouter:
    def __init__(self):
        self._extractors: Dict[str, object] = {
            ".pdf": PdfTextExtractor(),
            ".docx": DocxTextExtractor(),
            ".txt": TxtTextExtractor(),
        }

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".doc":
            raise UnsupportedDocumentError(
                "Khong ho tro file .doc trong V1. Vui long chuyen sang .docx truoc khi tai len."
            )
        extractor = self._extractors.get(suffix)
        if extractor is None:
            raise UnsupportedDocumentError(f"Unsupported file type: {suffix or 'unknown'}")
        return extractor.extract(path)
